# -*- coding: utf-8 -*-
"""
formkit — ابزارِ ساختِ قالب‌های Word فورم‌های رسمی (RTL، A4، آمادهٔ چاپ).

هر «توکن» به شکل {{Name}} و همیشه داخلِ یک ران تنها نوشته می‌شود، چون
DocxFormExport فقط جایگزینیِ درون‌رانی انجام می‌دهد و AssertNoTokensLeft
هر توکنِ شکسته یا جامانده را به خطا تبدیل می‌کند.
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from xml.sax.saxutils import escape
import os

W = nsdecls('w')

# ── پالت و اندازه‌ها ──────────────────────────────────────────────────────
NAVY      = "1F3864"
NAVY_SOFT = "2E5496"
LABEL_BG  = "EDF0F5"
BAND_BG   = "F5F7FA"
GRID      = "8C9BB0"
GRID_SOFT = "BFC8D6"
INK       = "1A1A1A"
MUTED     = "5A6472"

FONT_BODY  = "B Nazanin"
FONT_TITLE = "B Titr"
FONT_SYM   = "Segoe UI Symbol"

CELL_LINE = 200           # فاصلهٔ خطیِ دقیق داخلِ سلول‌ها (twips) — مهار ارتفاعِ ردیف
CONTENT_W = 10760          # عرضِ قابلِ استفادهٔ صفحه به twips (A4، حاشیهٔ ۱ سانتی)

TOKENS = set()


def reset_tokens():
    TOKENS.clear()


def tok(name):
    TOKENS.add(name)
    return "{{" + name + "}}"


# ── ران ──────────────────────────────────────────────────────────────────
def _rpr(bold, size, font, color, italic, underline):
    sz = int(round(size * 2))
    x = ['<w:rPr>']
    x.append('<w:rFonts w:ascii="%s" w:hAnsi="%s" w:cs="%s"/>' % (font, font, font))
    if bold:
        x.append('<w:b/><w:bCs/>')
    if italic:
        x.append('<w:i/><w:iCs/>')
    if color:
        x.append('<w:color w:val="%s"/>' % color)
    x.append('<w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (sz, sz))
    if underline:
        x.append('<w:u w:val="single"/>')
    x.append('<w:rtl/>')
    x.append('</w:rPr>')
    return "".join(x)


def _clean(text):
    """
    حذفِ اعرابِ اضافه‌کسره و «همزهٔ رویِ ه». فونت‌های B Nazanin/B Titr این دو
    را به‌شکلِ یک دایرهٔ جدا نشان می‌دهند و متنِ چاپی را کثیف می‌کنند.
    """
    return text.replace("ِ", "").replace("ٔ", "")


def run(p, text, bold=False, size=10, font=FONT_BODY, color=None,
        italic=False, underline=False):
    text = _clean(text)
    xml = ('<w:r %s>%s<w:t xml:space="preserve">%s</w:t></w:r>'
           % (W, _rpr(bold, size, font, color, italic, underline), escape(text)))
    p._p.append(parse_xml(xml))
    return p


# ── پاراگراف ─────────────────────────────────────────────────────────────
def _ppr(align, space_before, space_after, line, keep_next, border_bottom,
         shading, ind_right, ind_left, line_exact=None):
    x = ['<w:pPr>']
    if keep_next:
        x.append('<w:keepNext/>')
    x.append('<w:bidi/>')
    if shading:
        x.append('<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % shading)
    if border_bottom:
        x.append('<w:pBdr><w:bottom w:val="single" w:sz="%d" w:space="1" w:color="%s"/></w:pBdr>'
                 % (border_bottom[1], border_bottom[0]))
    sp = '<w:spacing w:before="%d" w:after="%d"' % (int(space_before * 20), int(space_after * 20))
    if line_exact:
        sp += ' w:line="%d" w:lineRule="exact"' % int(line_exact)
    elif line:
        sp += ' w:line="%d" w:lineRule="auto"' % int(line * 240)
    x.append(sp + '/>')
    if ind_right or ind_left:
        x.append('<w:ind w:right="%d" w:left="%d"/>' % (int(ind_right), int(ind_left)))
    x.append('<w:jc w:val="%s"/>' % align)
    x.append('</w:pPr>')
    return "".join(x)


def para(container, align="right", space_before=0, space_after=0, line=None,
         keep_next=False, border_bottom=None, shading=None, ind_right=0, ind_left=0,
         line_exact=None):
    p = container.add_paragraph()
    xml = '<w:pPr %s>%s' % (W, _ppr(align, space_before, space_after, line, keep_next,
                                    border_bottom, shading, ind_right, ind_left,
                                    line_exact)[len('<w:pPr>'):])
    p._p.insert(0, parse_xml(xml))
    return p


def text_para(container, text, bold=False, size=10, font=FONT_BODY, color=None,
              align="right", space_before=0, space_after=0, line=None,
              keep_next=False, border_bottom=None, shading=None):
    p = para(container, align, space_before, space_after, line, keep_next,
             border_bottom, shading)
    if text:
        run(p, text, bold=bold, size=size, font=font, color=color)
    return p


def page_break(container):
    """
    شکستِ صفحهٔ صریح. فقط جایی به کار می‌رود که فورم عمداً دوصفحه‌ای است و
    نمی‌خواهیم صفحهٔ دوم فقط بلوکِ امضاء باشد — با شکستِ دستی، مرزِ صفحه سرِ
    یک بخشِ کامل می‌افتد نه وسطِ آن.
    """
    p = para(container)
    p._p.append(parse_xml(
        '<w:r %s><w:br w:type="page"/></w:r>' % W))
    return p


def spacer(container, pts=4):
    p = para(container)
    pPr = p._p.find(qn('w:pPr'))
    sp = pPr.find(qn('w:spacing'))
    sp.set(qn('w:before'), "0")
    sp.set(qn('w:after'), "0")
    sp.set(qn('w:line'), str(int(pts * 20)))
    sp.set(qn('w:lineRule'), "exact")
    return p


def clear_cell(cell):
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)


# ── جدول ─────────────────────────────────────────────────────────────────
def _borders(color, sz, none):
    sides = ("top", "left", "bottom", "right", "insideH", "insideV")
    if none:
        body = "".join('<w:%s w:val="none" w:sz="0" w:space="0" w:color="auto"/>' % s
                       for s in sides)
    else:
        body = "".join('<w:%s w:val="single" w:sz="%d" w:space="0" w:color="%s"/>'
                       % (s, sz, color) for s in sides)
    return '<w:tblBorders>%s</w:tblBorders>' % body


def table(container, widths, borders=True, border_color=GRID, border_sz=4,
          cell_margin=54):
    """جدولِ RTL با عرضِ ثابتِ ستون‌ها (twips). ستونِ اول = سمتِ راست."""
    t = container.add_table(rows=0, cols=len(widths))
    tbl = t._tbl

    old = tbl.find(qn('w:tblPr'))
    if old is not None:
        tbl.remove(old)
    pr = ['<w:tblPr %s>' % W]
    pr.append('<w:tblW w:w="%d" w:type="dxa"/>' % sum(widths))
    pr.append('<w:jc w:val="right"/>')
    pr.append('<w:tblLayout w:type="fixed"/>')
    pr.append('<w:tblCellMar>')
    for side in ("top", "left", "bottom", "right"):
        m = 20 if side in ("top", "bottom") else cell_margin
        pr.append('<w:%s w:w="%d" w:type="dxa"/>' % (side, m))
    pr.append('</w:tblCellMar>')
    pr.append(_borders(border_color, border_sz, not borders))
    pr.append('<w:bidiVisual/>')
    pr.append('</w:tblPr>')
    tbl.insert(0, parse_xml("".join(pr)))

    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        tbl.remove(grid)
    g = ['<w:tblGrid %s>' % W]
    for w in widths:
        g.append('<w:gridCol w:w="%d"/>' % w)
    g.append('</w:tblGrid>')
    tbl.insert(1, parse_xml("".join(g)))

    t._widths = widths
    return t


def row(t, height=None):
    r = t.add_row()
    for i, c in enumerate(r.cells):
        clear_cell(c)
        set_width(c, t._widths[i])
    trPr = r._tr.get_or_add_trPr()
    trPr.append(parse_xml('<w:cantSplit %s/>' % W))
    if height:
        trPr = r._tr.get_or_add_trPr()
        h = OxmlElement('w:trHeight')
        h.set(qn('w:val'), str(int(height)))
        h.set(qn('w:hRule'), 'atLeast')
        trPr.append(h)
    return r


def set_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:tcW'))
    if old is not None:
        tcPr.remove(old)
    tcPr.insert(0, parse_xml('<w:tcW %s w:w="%d" w:type="dxa"/>' % (W, w)))


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:shd'))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (W, fill)))


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:vAlign'))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(parse_xml('<w:vAlign %s w:val="center"/>' % W))


def _ensure_p(cell):
    """هر <w:tc> باید دستِ‌کم یک بلوک داشته باشد وگرنه سند نامعتبر است."""
    if cell._tc.find(qn('w:p')) is None:
        cell._tc.append(parse_xml('<w:p %s/>' % W))


def merge(r, i, j):
    """ادغامِ افقیِ سلول‌های i تا j (اندیس از راست به چپ)."""
    for k in range(i, j + 1):
        _ensure_p(r.cells[k])
    m = r.cells[i].merge(r.cells[j])
    clear_cell(m)
    return m


def vmerge_column(t, col):
    """ادغامِ عمودیِ یک ستون در همهٔ ردیف‌های جدول."""
    cells = [t.rows[i].cells[col] for i in range(len(t.rows))]
    for c in cells:
        _ensure_p(c)
    m = cells[0]
    for c in cells[1:]:
        m = m.merge(c)
    clear_cell(m)
    return m


def finalize(doc):
    """پیش از ذخیره، سلول‌های خالی را یک پاراگرافِ تهی می‌دهیم."""
    for tc in doc.element.body.iter(qn('w:tc')):
        if tc.find(qn('w:p')) is None:
            tc.append(parse_xml('<w:p %s/>' % W))


# ── اجزای فورم ───────────────────────────────────────────────────────────
def label_cell(cell, text, size=8.5):
    shade(cell, LABEL_BG)
    vcenter(cell)
    p = para(cell, "right", line_exact=CELL_LINE)
    run(p, text, bold=True, size=size, color=INK)


def value_cell(cell, token_name=None, size=8.5, static=None, align="right", bold=False):
    vcenter(cell)
    p = para(cell, align, line_exact=CELL_LINE)
    if static:
        run(p, static, size=size, bold=bold)
    if token_name:
        run(p, tok(token_name), size=size, bold=bold)
    return p


def checkbox(p, token_name, caption, size=8, gap="  "):
    run(p, caption + " ", size=size)
    run(p, tok(token_name), size=size + 1, font=FONT_SYM)
    run(p, gap, size=size)


def choice_cell(cell, items, size=8, label=None):
    """سلولی که چند چک‌باکس در یک خط دارد."""
    vcenter(cell)
    p = para(cell, "right", line_exact=CELL_LINE)
    if label:
        run(p, label + " ", bold=True, size=size)
    for token_name, caption in items:
        checkbox(p, token_name, caption, size=size)
    return p


def section_bar(container, title, note=None, width=CONTENT_W, height=180):
    t = table(container, [width], borders=False, cell_margin=80)
    r = row(t, height=height)
    c = r.cells[0]
    shade(c, NAVY)
    vcenter(c)
    p = para(c, "right", space_before=2, space_after=2, keep_next=True)
    run(p, title, bold=True, size=9.5, font=FONT_TITLE, color="FFFFFF")
    if note:
        run(p, "    " + note, size=8.5, color="D6DCE8")
    sp = spacer(container, 3)
    sp._p.find(qn('w:pPr')).insert(0, parse_xml('<w:keepNext %s/>' % W))
    return t


def dotted(p, n=40, size=9.5):
    run(p, "." * n, size=size, color=MUTED)



# ── ردیفِ «برچسب | مقدار» ×۳ ──────────────────────────────────────────────
def val(name, size=8.5, static=None):
    return lambda c: value_cell(c, name, size=size, static=static)


def blank(static=None, size=8.5):
    return lambda c: value_cell(c, None, size=size, static=static)


def chks(items, size=8, label=None):
    return lambda c: choice_cell(c, items, size=size, label=label)


def kv_row(t, cells, height=1):
    """cells: فهرستِ (برچسب، سازنده)؛ سازنده روی سلولِ مقدار اجرا می‌شود."""
    r = row(t, height=height)
    for i, (cap, filler) in enumerate(cells):
        label_cell(r.cells[i * 2], cap)
        filler(r.cells[i * 2 + 1])
    return r


def six_widths(label=1430):
    w = [label, 0] * 3
    v = (CONTENT_W - label * 3) // 3
    for i in (1, 3, 5):
        w[i] = v
    w[5] = CONTENT_W - sum(w[:5])
    return w


# ── صفحه ─────────────────────────────────────────────────────────────────
def new_document(margins=(0.6, 0.5, 0.85, 0.85)):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT_BODY
    st.font.size = Pt(10)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:cs'), FONT_BODY)
    rf.set(qn('w:ascii'), FONT_BODY)
    rf.set(qn('w:hAnsi'), FONT_BODY)

    s = doc.sections[0]
    s.page_width = Cm(21)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(margins[0])
    s.bottom_margin = Cm(margins[1])
    s.left_margin = Cm(margins[2])
    s.right_margin = Cm(margins[3])
    s.footer_distance = Cm(0.5)
    s._sectPr.append(parse_xml('<w:bidi %s/>' % W))

    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    return doc


def _field(p, fld):
    for xml in (
        '<w:r %s><w:rPr><w:sz w:val="16"/><w:szCs w:val="16"/><w:color w:val="%s"/></w:rPr>'
        '<w:fldChar w:fldCharType="begin"/></w:r>' % (W, MUTED),
        '<w:r %s><w:instrText xml:space="preserve"> %s </w:instrText></w:r>' % (W, fld),
        '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % W,
        '<w:r %s><w:rPr><w:sz w:val="16"/><w:szCs w:val="16"/><w:color w:val="%s"/></w:rPr>'
        '<w:t>1</w:t></w:r>' % (W, MUTED),
        '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % W,
    ):
        p._p.append(parse_xml(xml))


def footer(doc, form_code):
    f = doc.sections[0].footer
    for p in list(f.paragraphs):
        p._p.getparent().remove(p._p)
    p = para(f, "center")
    run(p, form_code + "   •   کد پرونده: ", size=8, color=MUTED)
    run(p, tok("Code"), size=8, color=MUTED)
    run(p, "   •   تاریخ چاپ: ", size=8, color=MUTED)
    run(p, tok("PrintDate"), size=8, color=MUTED)
    run(p, "   •   صفحه ", size=8, color=MUTED)
    _field(p, "PAGE")
    run(p, " از ", size=8, color=MUTED)
    _field(p, "NUMPAGES")


def header_block(doc, title, right_rows, left_rows, subtitle=None, height=460):
    t = table(doc, [3080, 4600, 3080], borders=False, cell_margin=40)
    r = row(t, height=height)
    right, mid, left = r.cells[0], r.cells[1], r.cells[2]

    vcenter(right)
    for cap, name in right_rows:
        p = para(right, "right", space_after=1)
        run(p, cap + ": ", bold=True, size=9)
        run(p, tok(name), size=9)

    vcenter(mid)
    p = para(mid, "center")
    run(p, "به نام خدا", size=8.5, color=MUTED)
    p = para(mid, "center", space_after=1)
    run(p, tok("OrgName"), bold=True, size=10.5, font=FONT_TITLE, color=NAVY_SOFT)
    p = para(mid, "center", space_before=2)
    run(p, title, bold=True, size=14, font=FONT_TITLE, color=NAVY)
    if subtitle:
        p = para(mid, "center", space_before=1)
        run(p, subtitle, size=8.5, color=MUTED)

    vcenter(left)
    for cap, name in left_rows:
        p = para(left, "left", space_after=1)
        run(p, cap + ": ", bold=True, size=9)
        run(p, tok(name), size=9)

    para(doc, "right", space_before=1, space_after=4, border_bottom=(NAVY, 12))
    spacer(doc, 2)
    return t


def closing_block(doc, photo_caption, columns, height=1450, photo_w=3400):
    """
    بلوکِ پایانیِ فورم: یک کادرِ عکس (اختیاری) در راست و ستون‌های امضاء در چپ —
    همه در یک ردیف تا ارتفاعِ صفحه هدر نرود.
    """
    widths = [photo_w] if photo_caption else []
    rest = CONTENT_W - sum(widths)
    n = len(columns)
    w = rest // n
    widths += [w] * n
    widths[-1] = CONTENT_W - sum(widths[:-1])
    t = table(doc, widths, border_color=GRID_SOFT, border_sz=6)
    r = row(t, height=height)
    idx = 0
    if photo_caption:
        c = r.cells[0]
        shade(c, BAND_BG)
        text_para(c, photo_caption, bold=True, size=9, color=NAVY,
                  align="center", space_before=5)
        text_para(c, "(پس از چاپ الصاق و مهر گردد)", size=8, color=MUTED,
                  align="center", space_before=2)
        idx = 1
    for i, (caption, sub) in enumerate(columns):
        c = r.cells[idx + i]
        text_para(c, caption, bold=True, size=9, color=NAVY,
                  align="center", space_before=5)
        if sub:
            text_para(c, sub, size=8, color=MUTED, align="center", space_before=2)
    return t


def signature_block(doc, columns, height=1050):
    n = len(columns)
    w = CONTENT_W // n
    widths = [w] * n
    widths[-1] = CONTENT_W - w * (n - 1)
    t = table(doc, widths, border_color=GRID_SOFT, border_sz=6)
    r = row(t, height=height)
    for i, (caption, sub) in enumerate(columns):
        c = r.cells[i]
        p = para(c, "center", space_before=3, space_after=2)
        run(p, caption, bold=True, size=9.5, color=NAVY)
        if sub:
            p = para(c, "center")
            run(p, sub, size=8.5, color=MUTED)
    return t


def docs_checklist(doc, title, items, height=260):
    t = table(doc, [CONTENT_W], border_color=GRID_SOFT, border_sz=4, cell_margin=80)
    r = row(t, height=height)
    c = r.cells[0]
    shade(c, BAND_BG)
    vcenter(c)
    p = para(c, "right", space_before=2, space_after=2)
    run(p, title + " ", bold=True, size=9.5, color=NAVY)
    for token_name, caption in items:
        checkbox(p, token_name, caption, size=9, gap="     ")
    return t


def save(doc, out_dir, filename, manifest_path=None):
    finalize(doc)
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, filename)
    doc.save(path)
    names = sorted(TOKENS)
    if manifest_path:
        with open(manifest_path, "a", encoding="utf-8") as fh:
            fh.write("### " + filename + "  (" + str(len(names)) + " tokens)\n")
            fh.write("\n".join(names) + "\n\n")
    return path, names
